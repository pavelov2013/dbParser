import openpyxl as opx
import os.path
import docx
from datetime import datetime
import time

start_time = datetime.now()
wb = opx.load_workbook("res.xlsx")

sheet = wb.active
directory = 'referats/'

files = os.listdir(directory)

for el in files:
     sheet["A"+str(files.index(el)+1)].value = el


keys = list()
with open("keys.txt", encoding="utf-8") as file:
    keys = [l.strip() for l in file]



for i in files:
     iterator = 0
     document = docx.Document("referats/"+ i)
     doc = ""
     info = ""
     for el in document.paragraphs:
          doc += el.text
     doc = doc.lower()
     for el in keys:
          if el in doc:
               #print("ind = " + i + "\nkey = " +el + "\n\n")
               info += el+","
               iterator+=1
     sheet["Z"+str(files.index(i)+1)].value = info
     sheet["AD"+str(files.index(i)+1)].value = iterator
wb.save("res.xlsx")
print(datetime.now() - start_time)
